from datetime import datetime
import os
import tempfile
from tkinter import messagebox, filedialog
from pptx import Presentation
from pptx.util import Inches, Pt
from otgraficomodelo import OTGraficosModelo
from otgraficovista import OTGraficosVista

class OTGraficosControlador:
    def __init__(self, parent, repositorio_ot):
        self.repositorio = repositorio_ot
        self.modelo = OTGraficosModelo(self.repositorio)
        self.vista = OTGraficosVista(parent)
        self.vista.set_controlador(self)

        self.modo = "mensual"
        self.anio_actual = datetime.now().year

        self._cargar_combos()
        self._conectar_eventos()
        self.cargar_datos()

    def _cargar_combos(self):
        todas = self.modelo.obtener_todas()
        if not todas: return

        self.vista.combo_mes["values"] = [""] + list(range(1, 13))
        self.vista.combo_mes.set("")

        anios = sorted({ot["fecha"].year for ot in todas if ot.get("fecha")})
        if anios:
            self.vista.combo_anio["values"] = anios
            self.vista.combo_anio.set(self.anio_actual if self.anio_actual in anios else anios[0])

        self.vista.combo_planta["values"] = [""] + sorted({ot["planta"] for ot in todas if ot.get("planta")})
        self.vista.combo_linea["values"] = [""] + sorted({ot["linea"] for ot in todas if ot.get("linea")})
        
        self.vista.combo_tipo["values"] = ["", "CORRECTIVO INMEDIATO", "CORRECTIVO DIFERIDO", "MEJORA"]
        self.vista.combo_tipo.set("")

    def _conectar_eventos(self):
        combos = [self.vista.combo_mes, self.vista.combo_anio, self.vista.combo_planta, 
                  self.vista.combo_linea, self.vista.combo_tipo]
        for c in combos:
            c.bind("<<ComboboxSelected>>", lambda e: self.cargar_datos())

    def cambiar_modo(self, modo):
        self.modo = modo
        self.vista.combo_mes.configure(state="disabled" if modo == "anual" else "readonly")
        if modo == "anual": self.vista.combo_mes.set("")
        self.cargar_datos()

    def _obtener_filtros(self):
        m = self.vista.combo_mes.get()
        a = self.vista.combo_anio.get()
        return (
            self.vista.combo_planta.get() or None,
            self.vista.combo_linea.get() or None,
            int(m) if m else None,
            int(a) if a else self.anio_actual,
            self.vista.combo_tipo.get() or None
        )

    def cargar_datos(self):
        try:
            planta, linea, mes, anio, tipo_ot = self._obtener_filtros()
            if self.modo == "mensual":
                res = self.modelo.resumen(planta, linea, mes, anio, tipo_ot)
                self.vista.mostrar_grafico_resumen(res)
            else:
                res = self.modelo.resumen_anual(planta, linea, anio, tipo_ot)
                self.vista.mostrar_grafico_anual(res)
        except Exception as e:
            print(f"Error en gráficos: {e}")

    def generar_presentacion(self):
        """Crea la presentación en PowerPoint con ambos gráficos (Mensual y Anual) ocupando pantalla completa."""
        try:
            planta, linea, mes, anio, tipo_ot = self._obtener_filtros()
            planta_nombre = planta or self.modelo.obtener_planta_bd()
            
            # Nombre predeterminado del archivo
            partes_nombre = ["Reporte_Mantenimiento", planta_nombre]
            if mes:
                partes_nombre.append(f"Mes_{mes}")
            partes_nombre.append(str(anio))
            
            nombre_sugerido = "_".join(partes_nombre).replace(" ", "_") + ".pptx"

            filepath = filedialog.asksaveasfilename(
                defaultextension=".pptx",
                initialfile=nombre_sugerido,
                filetypes=[("Presentación de PowerPoint", "*.pptx")],
                title="Guardar Presentación Completa"
            )

            if not filepath:
                return

            # Obtener datos del modelo para ambos reportes
            res_mensual = self.modelo.resumen(planta, linea, mes, anio, tipo_ot)
            res_anual = self.modelo.resumen_anual(planta, linea, anio, tipo_ot)

            # Rutas temporales para las imágenes
            temp_dir = tempfile.gettempdir()
            img_mensual = os.path.join(temp_dir, "temp_grafico_mensual.png")
            img_anual = os.path.join(temp_dir, "temp_grafico_anual.png")

            # Exportar ambas imágenes a disco en 16:9
            self.vista.exportar_grafico_a_imagen("mensual", res_mensual, img_mensual)
            self.vista.exportar_grafico_a_imagen("anual", res_anual, img_anual)

            # Inicializar PowerPoint con formato 16:9 exacto
            prs = Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            blank_layout = prs.slide_layouts[6]  # Diapositiva limpia

            # Dimensiones para ocupar toda la filmina manteniendo margen mínimo
            left = Inches(0.3)
            top = Inches(0.3)
            width = Inches(12.733)
            height = Inches(6.9)

            # --- Diapositiva 1: Resumen Mensual ---
            slide1 = prs.slides.add_slide(blank_layout)
            slide1.shapes.add_picture(img_mensual, left, top, width=width, height=height)

            # --- Diapositiva 2: Evolución Anual ---
            slide2 = prs.slides.add_slide(blank_layout)
            slide2.shapes.add_picture(img_anual, left, top, width=width, height=height)

            # Guardar la presentación de PowerPoint
            prs.save(filepath)

            # Limpiar archivos temporales
            for img in (img_mensual, img_anual):
                if os.path.exists(img):
                    os.remove(img)

            messagebox.showinfo("Éxito", f"Presentación generada con ambos gráficos:\n{os.path.basename(filepath)}")

        except Exception as e:
            messagebox.showerror("Error", f"No se pudo generar la presentación:\n{e}")